#!/usr/bin/env python3
"""
Extract text from all docs in ./docs/ into a single JSON manifest.
Reads .docx, .pdf, .pptx. Output: tools/extracted_docs.json

Run: python3 tools/extract_docs.py
"""
import json
import os
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from pptx import Presentation


ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = ROOT / "docs"
OUT = ROOT / "tools" / "extracted_docs.json"


def extract_docx(path: Path) -> dict:
    doc = DocxDocument(path)
    sections = []
    current = {"heading": None, "paras": []}
    for p in doc.paragraphs:
        style_name = ""
        try:
            if p.style is not None and p.style.name:
                style_name = p.style.name
        except Exception:
            style_name = ""
        style = style_name.lower()
        text = p.text.strip()
        if not text:
            continue
        if style.startswith("heading"):
            if current["heading"] or current["paras"]:
                sections.append(current)
            current = {"heading": text, "paras": []}
        else:
            current["paras"].append(text)
    if current["heading"] or current["paras"]:
        sections.append(current)
    # Tables
    tables = []
    for tbl in doc.tables:
        rows = []
        for row in tbl.rows:
            rows.append([c.text.strip() for c in row.cells])
        if rows:
            tables.append(rows)
    return {"sections": sections, "tables": tables}


def extract_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append({"page": i + 1, "text": page.extract_text() or ""})
        except Exception as e:
            pages.append({"page": i + 1, "text": "", "error": str(e)})
    return {"pages": pages}


def extract_pptx(path: Path) -> dict:
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides):
        title = ""
        body = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if not t:
                        continue
                    if not title:
                        title = t
                    else:
                        body.append(t)
        slides.append({"slide": i + 1, "title": title, "body": body})
    return {"slides": slides}


def main():
    manifest = []
    for path in sorted(DOCS_DIR.rglob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        rel = path.relative_to(ROOT)
        entry = {
            "id": str(rel).replace("/", "__").replace(".", "_"),
            "path": str(rel),
            "filename": path.name,
            "category": path.parent.name if path.parent != DOCS_DIR else "General",
            "type": suffix.lstrip("."),
        }
        try:
            if suffix == ".docx":
                entry["content"] = extract_docx(path)
            elif suffix == ".pdf":
                entry["content"] = extract_pdf(path)
            elif suffix == ".pptx":
                entry["content"] = extract_pptx(path)
            else:
                continue
            manifest.append(entry)
            print(f"[ok] {rel}")
        except Exception as e:
            print(f"[err] {rel}: {e}")
            entry["error"] = str(e)
            manifest.append(entry)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(manifest, indent=2, ensure_ascii=False))
    print(f"\nWrote {OUT} ({len(manifest)} docs)")


if __name__ == "__main__":
    main()
